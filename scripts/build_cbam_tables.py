"""Build data/eu/cbam_tables.json from the official European Commission Excel files:
  - default_values_2026-08-06.xlsx  (DV correcting act, Annex I of Impl. Reg. 2025/2621 as amended by 2026/1740)
  - benchmarks_2026-02-06.xlsx      (CBAM benchmarks, Impl. Reg. 2025/2620; Column B = with CBAM factor applied)
Output: {"default_values": {country: {cn_digits: {"direct","indirect","total","route","desc"}}},
         "benchmarks_b": {cn8: {route or "": value}}, "sources": {...}}"""
import json
import re
from pathlib import Path

import openpyxl

EU = Path(__file__).resolve().parent.parent / "data" / "eu"


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def digits(v) -> str:
    return re.sub(r"\D", "", str(v or ""))


def default_values():
    wb = openpyxl.load_workbook(EU / "default_values_2026-08-06.xlsx", read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        if ws.title in ("Overview", "Version History"):
            continue
        rows = {}
        for row in ws.iter_rows(min_row=3, values_only=True):
            code, desc, d, i, t, route = (list(row) + [None] * 6)[:6]
            cd = digits(code)
            if len(cd) < 4:
                continue
            tot = num(t)
            if tot is None:
                continue
            rows[cd] = {"direct": num(d), "indirect": num(i), "total": tot,
                        "route": (str(route or "").strip("() ") or ""), "desc": str(desc or "")[:160]}
        out[ws.title] = rows
    return out


def benchmarks():
    wb = openpyxl.load_workbook(EU / "benchmarks_2026-02-06.xlsx", read_only=True, data_only=True)
    ws = wb["Benchmarks"]
    out, cur = {}, None
    for row in ws.iter_rows(min_row=2, values_only=True):
        code, desc, a, ra, b, rb = (list(row) + [None] * 6)[:6]
        cd = digits(code)
        if len(cd) >= 4:
            cur = cd
        if cur is None or num(b) is None:
            continue
        out.setdefault(cur, {})[str(rb or "").strip("() ")] = num(b)
    return out


if __name__ == "__main__":
    dv, bm = default_values(), benchmarks()
    data = {"default_values": dv, "benchmarks_b": bm, "sources": {
        "default_values": "European Commission, DV correcting act final update 06.08.2026 (Annex I, Impl. Reg. 2025/2621 as amended by 2026/1740)",
        "benchmarks": "European Commission, CBAM Benchmarks 2026-02-06 (Impl. Reg. 2025/2620), Column B",
        "url": "https://taxation-customs.ec.europa.eu/carbon-border-adjustment-mechanism/cbam-legislation-and-guidance_en"}}
    (EU / "cbam_tables.json").write_text(json.dumps(data, ensure_ascii=False))
    print(len(dv), "countries;", sum(len(v) for v in dv.values()), "DV rows;", len(bm), "benchmark codes")
    for c, code in (("India", "72071114"), ("China", "72071114"), ("Indonesia", "72071114"), ("Turkey", "72071114"), ("Uzbekistan", "72071114")):
        print(c, dv.get(c, {}).get(code))
    print("BM 72071114", bm.get("72071114"), "BM 7208", {k: v for k, v in bm.items() if k.startswith("7208")} and list({k: v for k, v in bm.items() if k.startswith('7208')}.items())[:2])
